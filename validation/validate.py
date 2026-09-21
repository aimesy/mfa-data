"""Reproduce release integrity, evidence, scope and disclosure checks.

Run from any directory with the pinned requirements installed. Source interpretation
and actual visual review are recorded separately; this script cannot perform them.
"""
import csv,hashlib,io,json,pathlib,re,sys,zipfile
from decimal import Decimal
import pymupdf
from pypdf import PdfReader
from openpyxl import load_workbook
import pdf_graph as graph

ROOT=pathlib.Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads((ROOT/p).read_text(encoding='utf-8'))
def local(p):
 assert isinstance(p,str) and not re.match(r'^[A-Za-z]:|^/|^\\|^https?:',p),p
 target=(ROOT/p.split('#')[0]).resolve();assert target.is_relative_to(ROOT) and target.is_file(),p
 return target

def validate_data(rows):
 expected={'fremont-2020-21-parkland':('7398609.00',4,'Parkland Fee'), 'fremont-2020-21-park-facilities':('7856947.00',6,'Park Facilities Fee')}
 assert len(rows)==len(expected) and {r['record_id'] for r in rows}==set(expected)
 for r in rows:
  amount,page,program=expected[r['record_id']]
  assert Decimal(r['value_usd'])==Decimal(amount)
  assert Decimal(r['source_value_text'].replace('$','').replace(',',''))==Decimal(amount)
  assert r['physical_pdf_page']==page and r['printed_page']==str(page) and r['fee_program']==program
  assert r['legal_entity']=='City of Fremont' and r['fiscal_year']=='2020-21'
  assert r['publication_title']=='Development Impact Fee Annual Report for Period Ending June 30, 2021'
  assert (r['fiscal_year_start'],r['fiscal_year_end'])==('2020-07-01','2021-06-30')
  assert r['fee_category']=='parks_recreation_open_space' and r['land_use']=='residential'
  assert r['measure_type']=='reported_residential_fee_collections'
  assert r['validation_status']=='accepted_reported_collection_scope' and r['cash_status']=='not_established'
  assert r['gross_or_net']=='not_established' and r['accounting_basis']=='not_stated_in_annual_report'
  assert r['derived_operands'] is None
  assert r['source_sha256']==sha(local(r['source_pdf']))
  for k in ['extract_pdf','highlight_pdf','attribution_evidence']:local(r[k])
  with pymupdf.open(local(r['source_pdf'])) as d:
   t=d[page-1].get_text();assert r['source_value_text'] in t and r['source_label'] in t
  ev=next(e for e in read('evidence/index.json') if e['source_id']==r['source_id'] and e['physical_pdf_page']==page)
  assert (r['extract_pdf'],r['highlight_pdf'])==(ev['extract_pdf'],ev['highlight_pdf'])
  assert r['attribution_evidence']=='evidence/extracts/fremont-2020-21-p017.pdf'
 return len(rows)

def privacy(files):
 patterns=[r'[A-Za-z]:[\\/]+Users[\\/]',r'/ho'+r'me/[^/]+/',r'account-[0-9]+',r'-----BEGIN [A-Z ]*PRIVATE KEY-----',r'gh[pousr]_[A-Za-z0-9]{20,}',r'sk-[A-Za-z0-9_-]{20,}',r'"(?:threadId|session_id|accountFingerprint)"\s*:']
 rx=[re.compile(p,re.I) for p in patterns];scanned=0;members=0;pdf_objects=0
 def scan(b,label):
  nonlocal scanned
  scanned+=1;t=b.decode('utf-8',errors='replace')
  assert not any(p.search(t) for p in rx),'Private operational content: '+label
 for p in files:
  rel=p.relative_to(ROOT).as_posix();assert p.suffix.lower() not in ['.sqlite','.log','.jsonl','.exe','.key','.pem'],rel
  raw=p.read_bytes();scan(raw,rel)
  if zipfile.is_zipfile(p):
   with zipfile.ZipFile(p) as z:
    for item in z.infolist():
     assert item.file_size<20_000_000 and not item.filename.startswith(('/', '\\')) and '..' not in pathlib.PurePosixPath(item.filename).parts
     scan(z.read(item),rel+':'+item.filename);members+=1
  if p.suffix.lower()=='.pdf':
   reader=PdfReader(io.BytesIO(raw))
   active={n for generation,entries in reader.xref.items() if generation!=65535 for n in entries if n>0 and not reader.xref_free_entry.get(generation,{}).get(n,False)}|set(reader.xref_objStm)
   with pymupdf.open(p) as d:
    assert d.embfile_count()==0,'Embedded attachment requires separate disclosure review'
    for i in sorted(active):
     scan(d.xref_object(i,compressed=False).encode(),rel+':object');pdf_objects+=1
     if d.xref_is_stream(i):
      stream=d.xref_stream(i);assert len(stream)<=75_000_000;scan(stream,rel+':stream')
 return {'passed':True,'scanned_payloads':scanned,'archive_members':members,'pdf_objects':pdf_objects}

def main():
 manifest=read('manifest.json');files=[local(x['path']) for x in manifest['files']]
 actual={p.relative_to(ROOT).as_posix() for p in ROOT.rglob('*') if p.is_file() and '.git' not in p.relative_to(ROOT).parts and '__pycache__' not in p.parts and p.name!='manifest.json'}
 assert actual=={x['path'] for x in manifest['files']},'Manifest file set differs'
 for x,p in zip(manifest['files'],files):assert p.stat().st_size==x['bytes'] and sha(p)==x['sha256'],x['path']
 rows=read('data/reported-fee-collections.json');n=validate_data(rows)
 for name in ['residential-cash-receipts','capital-spending','residential-funded-shares']:
  with (ROOT/'data'/f'{name}.csv').open(newline='',encoding='utf-8') as f:assert list(csv.DictReader(f))==[]
 with (ROOT/'data/reported-fee-collections.csv').open(newline='',encoding='utf-8') as f:
  csvrows=list(csv.DictReader(f))
  assert [{k:'' if v is None else str(v) for k,v in r.items()} for r in rows]==csvrows
 wb=load_workbook(ROOT/'data/mfa-preliminary.xlsx',read_only=True,data_only=False)
 ws=wb['Reported collections'];values=list(ws.values);assert list(values[0])==list(rows[0]) and len(values)==n+1
 for row,vals in zip(rows,values[1:]):
  for key,value in zip(values[0],vals):
   if key=='value_usd':assert Decimal(str(value))==Decimal(row[key])
   else:assert value==row[key],key
 for sheet,csvname in [('Residential cash','residential-cash-receipts'),('Capital spending','capital-spending'),('Funded shares','residential-funded-shares')]:
  assert wb[sheet].max_row==1
  with (ROOT/'data'/f'{csvname}.csv').open(newline='',encoding='utf-8') as f:assert list(next(wb[sheet].values))==next(csv.reader(f))
 index=read('sources/index.json');assert len(index)==3 and sum(p['selected_data_rows'] for p in index)==n
 values=list(wb['Source index'].values);assert list(values[0])==list(index[0])
 assert [list(p.values()) for p in index]==[list(v) for v in values[1:]]
 with (ROOT/'sources/index.csv').open(newline='',encoding='utf-8') as f:assert list(csv.DictReader(f))==[{k:'' if v is None else str(v) for k,v in p.items()} for p in index]
 for p in index:
  f=local(p['original_pdf']);assert sha(f)==p['source_sha256'] and f.stat().st_size==p['source_bytes']
  with pymupdf.open(f) as d:assert len(d)==p['physical_pages']
 for e in read('evidence/index.json'):
  s,x,h=(local(e[k]) for k in ['source_pdf','extract_pdf','highlight_pdf']);page=e['physical_pdf_page']-1
  for p,k in [(s,'source_sha256'),(x,'extract_sha256'),(h,'highlight_sha256')]:assert sha(p)==e[k]
  sg,_=graph.joint_snapshot(s.read_bytes(),page);xg,_=graph.joint_snapshot(x.read_bytes(),0)
  hg,_=graph.joint_snapshot(h.read_bytes(),0,annotation_limit=e['original_annotation_count'],original_annots_present=e['original_annots_present'])
  assert sg==xg==hg
  assert graph.annotation_count(h.read_bytes(),0)==e['original_annotation_count']+e['added_annotations']
  pymupdf.TOOLS.mupdf_warnings(reset=True)
  with pymupdf.open(s) as sd,pymupdf.open(x) as xd,pymupdf.open(h) as hd:
   assert sd[page].get_text()==xd[0].get_text()==hd[0].get_text()
   a=sd[page].get_pixmap(dpi=144,alpha=False);b=xd[0].get_pixmap(dpi=144,alpha=False)
   assert (a.width,a.height,a.samples)==(b.width,b.height,b.samples)
   original=hd[0].get_pixmap(dpi=144,alpha=False,annots=False);unmarked=xd[0].get_pixmap(dpi=144,alpha=False,annots=False)
   assert original.samples==unmarked.samples
   hp=hd[0];anns=list(hp.annots() or []);squares=[a for a in anns if a.type[1]=='Square'];assert len(squares)==e['added_annotations']
   rects=[r for t in e['reviewed_text'] for r in t['rectangles']]
   for a,r in zip(squares,rects):
    assert all(abs(x-y)<0.001 for x,y in zip(a.rect,r))
    assert a.colors['fill'] in (None,[]) and a.colors['stroke'] is not None
    assert hd.xref_get_key(a.xref,'IC')[0]=='null'
   for t in e['reviewed_text']:assert t['text'] in xd[0].get_text()
  assert not pymupdf.TOOLS.mupdf_warnings(reset=True)
 # Check Markdown file links; fragments do not change the underlying file target.
 links=0
 for p in files:
  if p.suffix=='.md':
   for target in re.findall(r'\]\(([^)]+)\)',p.read_text(encoding='utf-8')):
    if target.startswith(('http://','https://','#')):continue
    q=(p.parent/target.split('#')[0]).resolve();assert q.is_relative_to(ROOT) and q.is_file(),target;links+=1
 result={'passed':True,'manifest_files':len(files),'source_reports':len(index),'reported_collection_rows':n,'verified_cash_rows':0,'funded_shares':0,'exact_evidence_pairs':len(read('evidence/index.json')),'markdown_file_links':links,'privacy':privacy(files+[ROOT/'manifest.json']),'semantic_scope':'reported collections only; no cash or funded-share acceptance'}
 print(json.dumps(result,indent=2))
if __name__=='__main__':main()
